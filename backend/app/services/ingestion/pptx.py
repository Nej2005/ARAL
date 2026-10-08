"""PPTX -> slides: title first, then text frames in reading order (BACKEND.md §5.2)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from app.services.ingestion import ExtractedPage


def _iter_shapes(shapes):
    """Flatten group shapes so every text frame / table is visited."""
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def _frame_lines(text_frame) -> list[str]:
    lines = []
    for para in text_frame.paragraphs:
        txt = "".join(run.text for run in para.runs).rstrip()
        if not txt.strip():
            continue
        level = para.level or 0
        lines.append("  " * level + "• " + txt.strip())
    return lines


def _table_lines(table) -> list[str]:
    lines = []
    for row in table.rows:
        cells = [" ".join(c.text.split()) for c in row.cells]
        if any(cells):
            lines.append(" | ".join(cells))
    return lines


def _pos(sh) -> tuple[int, int]:
    top = int(sh.top) if sh.top is not None else 0
    left = int(sh.left) if sh.left is not None else 0
    return top, left


def extract_pptx(path: str | Path) -> list[ExtractedPage]:
    prs = Presentation(str(path))
    out: list[ExtractedPage] = []
    for n, slide in enumerate(prs.slides, start=1):
        title = None
        title_shape_id = None
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                t = " ".join(slide.shapes.title.text_frame.text.split())
                if t:
                    title = t
                    title_shape_id = slide.shapes.title.shape_id
        except Exception:  # some layouts raise on .title
            pass

        blocks: list[tuple[tuple[int, int], list[str]]] = []
        for sh in _iter_shapes(slide.shapes):
            if sh.shape_id == title_shape_id:
                continue
            if sh.has_text_frame:
                lines = _frame_lines(sh.text_frame)
                if lines:
                    blocks.append((_pos(sh), lines))
            elif getattr(sh, "has_table", False) and sh.has_table:
                lines = _table_lines(sh.table)
                if lines:
                    blocks.append((_pos(sh), lines))
        # Reading order: top -> bottom, then left -> right (rounded so near-equal tops tie).
        blocks.sort(key=lambda b: (round(b[0][0] / Emu(200000)), b[0][1]))

        lines: list[str] = []
        if title:
            lines.append(title)
        for _, blk in blocks:
            lines.extend(blk)
        out.append(ExtractedPage(page_no=n, title=title, text="\n".join(lines)))
    return out
