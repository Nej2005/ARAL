"""Builds the sample lesson files used by the tests (deterministic content)."""

from __future__ import annotations

from pathlib import Path

from fpdf import FPDF
from pptx import Presentation
from pptx.util import Inches, Pt

FIXTURE_DIR = Path(__file__).parent

# Definitions are "Term – meaning." ; facts are plain sentences starting with "The ".
BIO_SLIDES = [
    ("Photosynthesis", [
        "Photosynthesis – the process by which green plants use sunlight to synthesize food from carbon dioxide and water.",
        "Chlorophyll – the green pigment in chloroplasts that absorbs light energy, mostly in the blue and red wavelengths.",
        "The chloroplast is the organelle where photosynthesis takes place.",
    ]),
    ("Photosynthesis (cont.)", [
        "Stomata – small pores on the underside of a leaf that allow carbon dioxide to enter and oxygen to leave.",
        "Transpiration – the loss of water vapor through the stomata of leaves.",
        "The light-dependent reactions take place in the thylakoid membranes of the chloroplast.",
    ]),
    ("Calvin Cycle", [
        "Calvin cycle – the series of reactions in the stroma that use ATP and NADPH to convert carbon dioxide into glucose.",
        "Rubisco – the enzyme that fixes carbon dioxide during the Calvin cycle.",
        "RuBP – a five-carbon sugar that combines with carbon dioxide at the start of the Calvin cycle.",
        "The Calvin cycle produces one molecule of G3P for every three molecules of carbon dioxide fixed.",
    ]),
    ("Thank you!", ["Questions?"]),
]

RESP_PAGES = [
    ("Cellular Respiration", [
        "Cellular respiration – the process of breaking down glucose to release the energy stored in its bonds.",
        "ATP – the main energy-carrying molecule of the cell, which releases energy when its third phosphate group is removed.",
        "The mitochondria is known as the powerhouse of the cell.",
    ]),
    ("Glycolysis", [
        "Glycolysis – the first stage of cellular respiration, which breaks one glucose molecule into two molecules of pyruvate.",
        "Pyruvate – the three-carbon molecule produced by glycolysis.",
        "The glycolysis stage takes place in the cytoplasm and produces a net gain of two ATP.",
    ]),
    ("Krebs Cycle", [
        "Krebs cycle – a series of reactions in the mitochondrial matrix that releases carbon dioxide and produces NADH and FADH2.",
        "NADH – an electron carrier that brings high-energy electrons to the electron transport chain.",
        "The Krebs cycle releases two molecules of carbon dioxide for every acetyl-CoA that enters it.",
    ]),
    ("Electron Transport", [
        "Oxygen – the final electron acceptor in the electron transport chain, which combines with electrons and hydrogen ions to form water.",
        "Fermentation – an anaerobic process that produces ATP without oxygen.",
        "The fermentation process produces only two ATP per glucose molecule.",
    ]),
]


def build_pptx(path: Path, slides=BIO_SLIDES, variant: str | None = None) -> Path:
    prs = Presentation()
    if variant:
        prs.core_properties.title = variant  # different bytes, identical text
    layout = prs.slide_layouts[1]  # title + content
    for title, lines in slides:
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        body = s.placeholders[1].text_frame
        body.text = lines[0]
        for line in lines[1:]:
            p = body.add_paragraph()
            p.text = line
    # A table slide: term | meaning
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Key Terms"
    rows = [("Carotenoids", "accessory pigments that absorb light in the blue-green range and give leaves yellow and orange colors"),
            ("Thylakoid", "a membrane-bound sac inside the chloroplast where the light-dependent reactions occur")]
    shape = s.shapes.add_table(len(rows), 2, Inches(0.5), Inches(1.5), Inches(9), Inches(1.5))
    for i, (a, b) in enumerate(rows):
        shape.table.cell(i, 0).text = a
        shape.table.cell(i, 1).text = b
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def build_pdf(path: Path, pages=RESP_PAGES, header: str | None = "BIO 101 – Lesson 4", footer: bool = True) -> Path:
    from app.services.export import FONT_DIR

    pdf = FPDF()
    pdf.add_font("Noto", "", FONT_DIR / "NotoSans-Regular.ttf")
    pdf.add_font("Noto", "B", FONT_DIR / "NotoSans-Bold.ttf")
    pdf.set_auto_page_break(auto=False)
    for i, (title, lines) in enumerate(pages, start=1):
        pdf.add_page()
        pdf.set_font("Noto", size=11)
        if header:
            pdf.cell(0, 8, header, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
        pdf.set_font("Noto", "B", 14)
        pdf.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Noto", size=11)
        for line in lines:
            pdf.multi_cell(0, 7, line, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        if footer:
            pdf.set_y(-20)
            pdf.cell(0, 8, f"Page {i}", align="C", new_x="LMARGIN", new_y="NEXT")
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))
    return path


def build_empty_pdf(path: Path) -> Path:
    pdf = FPDF()
    pdf.add_page()
    pdf.output(str(path))
    return path
